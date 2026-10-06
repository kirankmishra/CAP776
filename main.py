import openpyxl

def load_data(file_name):
    records = []

    try:
        workbook = openpyxl.load_workbook(file_name, data_only=True)

        if "Daily Log" not in workbook.sheetnames:
            print("Daily Log sheet not found.")
            return []

        sheet = workbook["Daily Log"]

        for row in range(6, sheet.max_row + 1):
            date = sheet.cell(row, 1).value
            if date is None:
                continue

            notes = sheet.cell(row, 14).value
            if notes and "Example row" in str(notes):
                continue

            try:
                sleep = float(sheet.cell(row, 2).value or 0)
                fitness = float(sheet.cell(row, 3).value or 0)
                study = float(sheet.cell(row, 4).value or 0)
                coding = float(sheet.cell(row, 5).value or 0)
                class_time = float(sheet.cell(row, 6).value or 0)
                classes = float(sheet.cell(row, 7).value or 0)
                other = float(sheet.cell(row, 8).value or 0)

                feeling = str(sheet.cell(row, 11).value or "okay").strip().lower()
                satisfaction = str(sheet.cell(row, 12).value or "okay").strip().lower()
                energy = str(sheet.cell(row, 13).value or "medium").strip().lower()

                if min(sleep, fitness, study, coding, class_time, other) < 0:
                    continue

                total = sleep + fitness + study + coding + class_time + other
                if total > 1440:
                    continue

                free_time = 1440 - total

                records.append({
                    "date": date,
                    "sleep": sleep,
                    "fitness": fitness,
                    "study": study,
                    "coding": coding,
                    "class": class_time,
                    "classes": classes,
                    "other": other,
                    "total": total,
                    "free": free_time,
                    "feeling": feeling,
                    "satisfaction": satisfaction,
                    "energy": energy
                })

            except (ValueError, TypeError):
                continue

        return records

    except FileNotFoundError:
        print("Excel file not found.")
        return []
    except Exception as error:
        print("Error:", error)
        return []


def average(records, key):
    if len(records) == 0:
        return 0
    return sum(r[key] for r in records) / len(records)


def calculate_tpi(records):
    return average(records, "coding")


def calculate_aai(records):
    if not records:
        return 0
    total = sum(r["study"] + r["class"] for r in records)
    return total / len(records)


def calculate_phai(records):
    return average(records, "fitness")


def calculate_sri(records):
    return average(records, "sleep")


def calculate_abi(records):
    return average(records, "free")


def calculate_tui(records):
    return average(records, "total")


def calculate_ei(records):
    feeling = {"excellent": 5, "good": 4, "okay": 3, "neutral": 3, "low": 2, "stressed": 1}
    satisfaction = {"very satisfied": 5, "satisfied": 4, "okay": 3, "neutral": 3, "unsatisfied": 2, "very unsatisfied": 1}
    energy = {"high": 5, "medium": 3, "low": 1}

    total = 0
    for r in records:
        total += feeling.get(r["feeling"], 3)
        total += satisfaction.get(r["satisfaction"], 3)
        total += energy.get(r["energy"], 3)

    return total / (3 * len(records))


def calculate_dci(records):
    from datetime import date
    # Kiran's activity recording period: 15 Sept – 4 Oct 2026
    start = date(2026, 9, 15)
    end = date(2026, 10, 4)

    dates = set()
    for r in records:
        current = r["date"]
        if hasattr(current, "date"):
            current = current.date()
        if start <= current <= end:
            dates.add(current)

    # Expected days = 20
    return (len(dates) / 20) * 100


def calculate_pai(tpi, aai, phai, sri, tui, ei, dci):
    return (0.15 * tpi + 0.20 * aai + 0.15 * phai +
            0.20 * sri + 0.15 * tui + 0.10 * ei + 0.05 * dci)


def sleep_energy(records):
    print("\n1. Sleep vs Energy")
    for level in ["high", "medium", "low"]:
        values = [r["sleep"] for r in records if r["energy"] == level]
        if values:
            avg = sum(values) / len(values)
            print(level.title(), "-", len(values), "days, average sleep:", round(avg, 1), "min")


def study_satisfaction(records):
    print("\n2. Study + Class vs Satisfaction")
    for level in ["very satisfied", "satisfied", "okay", "neutral", "unsatisfied", "very unsatisfied"]:
        values = [r["study"] + r["class"] for r in records if r["satisfaction"] == level]
        if values:
            avg = sum(values) / len(values)
            print(level.title(), "-", len(values), "days, average academic time:", round(avg, 1), "min")


def coding_energy(records):
    print("\n3. Coding vs Energy")
    for level in ["high", "medium", "low"]:
        values = [r["coding"] for r in records if r["energy"] == level]
        if values:
            avg = sum(values) / len(values)
            print(level.title(), "-", len(values), "days, average coding:", round(avg, 1), "min")


# ------------------------------------------------------------
# MAIN PROGRAM
# ------------------------------------------------------------

file_name = "12628971.xlsx"

print("=" * 55)
print("           MY DATA - MY STORY")
print("       PERSONAL ACTIVITY ANALYSIS")
print("=" * 55)

data = load_data(file_name)

if data:
    print("\nExcel data loaded successfully!")
    print("Valid recorded rows:", len(data))

    tpi = calculate_tpi(data)
    aai = calculate_aai(data)
    phai = calculate_phai(data)
    sri = calculate_sri(data)
    abi = calculate_abi(data)
    tui = calculate_tui(data)
    ei = calculate_ei(data)
    dci = calculate_dci(data)

    pai = calculate_pai(tpi, aai, phai, sri, tui, ei, dci)

    print("\n" + "=" * 55)
    print("              ACTIVITY INDICES")
    print("=" * 55)
    print("Valid Recorded Days :", len(data))
    print("TPI  :", round(tpi, 2), "min/day")
    print("AAI  :", round(aai, 2), "min/day")
    print("PhAI :", round(phai, 2), "min/day")
    print("SRI  :", round(sri, 2), "min/day")
    print("ABI  :", round(abi, 2), "min/day")
    print("TUI  :", round(tui, 2), "min/day")
    print("EI   :", round(ei, 2), "/ 5")
    print("DCI  :", round(dci, 2), "%")
    print("-" * 55)
    print("PERSONAL ACTIVITY INDEX (PAI):", round(pai, 2))
    print("-" * 55)

    print("\n" + "=" * 55)
    print("           RELATIONSHIP ANALYSIS")
    print("=" * 55)
    sleep_energy(data)
    study_satisfaction(data)
    coding_energy(data)

    print("\n" + "=" * 55)
    print("Analysis completed successfully.")
    print("=" * 55)

else:
    print("No valid data found.")